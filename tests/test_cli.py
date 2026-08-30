"""Test della CLI: exit code, formati di output, batch e gestione errori."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from typer.testing import CliRunner

from nvdlocal.cli import app, read_inventory
from nvdlocal.db import connect, init_db, upsert_cves

from .conftest import FIXTURE_FILES, vulnerabilities

runner = CliRunner()


@pytest.fixture
def db_file(tmp_path: Path) -> Path:
    """Database su file popolato con le fixture, per invocare la CLI."""
    path = tmp_path / "nvd.db"
    conn = connect(path, create=True)
    init_db(conn)
    for name in FIXTURE_FILES:
        upsert_cves(conn, vulnerabilities(name))
    with conn:
        conn.execute(
            "INSERT INTO kev (cve_id, date_added, ransomware, due_date) VALUES (?, ?, ?, ?)",
            ("CVE-2021-41773", "2021-11-03", "Known", "2021-11-17"),
        )
        conn.execute(
            "INSERT INTO epss (cve_id, score, percentile, updated) VALUES (?, ?, ?, ?)",
            ("CVE-2021-41773", 0.97542, 0.99981, "2026-08-29"),
        )
        # Voce sintetica (identificatori volutamente inesistenti) per poter
        # verificare l'exit code su un risultato di severita' sotto HIGH:
        # nessuna delle fixture reali ha MEDIUM come metrica migliore.
        conn.execute(
            "INSERT INTO cve (cve_id, published, vuln_status, description, "
            "cvss_version, cvss_score, cvss_severity) "
            "VALUES ('CVE-1900-0001', '1900-01-01T00:00:00.000', 'Analyzed', "
            "'voce sintetica di test', '3.1', 5.3, 'MEDIUM')"
        )
        conn.execute(
            "INSERT INTO cpe_match (cve_id, config_index, node_index, negate, "
            "vulnerable, criteria, part, vendor, product, version) "
            "VALUES ('CVE-1900-0001', 0, 0, 0, 1, "
            "'cpe:2.3:a:esempio:prodotto_test:1.0:*:*:*:*:*:*:*', 'a', "
            "'esempio', 'prodotto_test', '1.0')"
        )
    conn.close()
    return path


def _run(db_file: Path, *args: str):
    return runner.invoke(app, ["--db", str(db_file), *args])


# --------------------------------------------------------------------------- #
# Exit code di search
# --------------------------------------------------------------------------- #
def test_search_exit_code_1_on_high_severity(db_file: Path) -> None:
    result = _run(db_file, "search", "--product", "httpd", "--version", "2.4.49")
    assert result.exit_code == 1
    assert "CVE-2021-41773" in result.stdout


def test_search_exit_code_0_when_nothing_found(db_file: Path) -> None:
    result = _run(db_file, "search", "--product", "httpd", "--version", "1.3.0")
    assert result.exit_code == 0
    assert "Nessuna CVE trovata" in result.stdout


def test_search_exit_code_0_when_below_high(db_file: Path) -> None:
    """Trova un risultato MEDIUM: ci sono CVE, ma nessuna >= HIGH -> exit code 0."""
    result = _run(
        db_file, "search", "--product", "prodotto_test", "--version", "1.0"
    )
    assert result.exit_code == 0
    assert "CVE-1900-0001" in result.stdout
    assert "MEDIUM" in result.stdout


def test_search_exit_code_0_when_filters_remove_everything(db_file: Path) -> None:
    result = _run(
        db_file,
        "search",
        "--product",
        "httpd",
        "--version",
        "2.4.49",
        "--min-score",
        "10.0",
    )
    assert result.exit_code == 0
    assert "Nessuna CVE trovata" in result.stdout


def test_search_rejects_meaningless_version(db_file: Path) -> None:
    result = _run(db_file, "search", "--product", "httpd", "--version", "*")
    assert result.exit_code == 2
    assert "versione non valida" in result.output


# --------------------------------------------------------------------------- #
# Formati di output
# --------------------------------------------------------------------------- #
def test_search_json_format(db_file: Path) -> None:
    result = _run(
        db_file, "search", "--product", "httpd", "--version", "2.4.49", "--format", "json"
    )
    payload = json.loads(result.stdout)
    assert payload["count"] >= 1
    first = payload["results"][0]
    for field in (
        "cve_id",
        "cvss_score",
        "cvss_severity",
        "cvss_vector",
        "cwe",
        "kev",
        "epss_score",
        "published",
        "match_type",
        "criteria",
        "version_range",
    ):
        assert field in first


def test_search_csv_to_file(db_file: Path, tmp_path: Path) -> None:
    output = tmp_path / "out.csv"
    result = _run(
        db_file,
        "search",
        "--product",
        "httpd",
        "--version",
        "2.4.49",
        "--format",
        "csv",
        "--output",
        str(output),
    )
    assert result.exit_code == 1
    rows = list(csv.DictReader(output.open(encoding="utf-8")))
    assert any(row["cve_id"] == "CVE-2021-41773" for row in rows)
    assert rows[0]["kev"] in ("yes", "no")


def test_search_by_cpe_takes_version_from_cpe(db_file: Path) -> None:
    result = _run(
        db_file,
        "search",
        "--cpe",
        "cpe:2.3:a:apache:http_server:2.4.49:*:*:*:*:*:*:*",
        "--format",
        "json",
    )
    payload = json.loads(result.stdout)
    assert "CVE-2021-41773" in {r["cve_id"] for r in payload["results"]}


def test_no_include_all_versions(db_file: Path) -> None:
    result = _run(
        db_file,
        "search",
        "--product",
        "vsftpd",
        "--version",
        "3.0.5",
        "--no-include-all-versions",
    )
    assert result.exit_code == 0
    assert "Nessuna CVE trovata" in result.stdout


def test_only_kev(db_file: Path) -> None:
    result = _run(
        db_file,
        "search",
        "--product",
        "httpd",
        "--version",
        "2.4.49",
        "--only-kev",
        "--format",
        "json",
    )
    payload = json.loads(result.stdout)
    assert {r["cve_id"] for r in payload["results"]} == {"CVE-2021-41773"}


# --------------------------------------------------------------------------- #
# Errori e ambiguita'
# --------------------------------------------------------------------------- #
def test_missing_database_is_explained(tmp_path: Path) -> None:
    result = runner.invoke(
        app,
        ["--db", str(tmp_path / "assente.db"), "search", "-p", "httpd", "-V", "2.4.49"],
    )
    assert result.exit_code == 2
    assert "sync --full" in result.output


def test_ambiguous_product_lists_candidates(db_file: Path) -> None:
    result = _run(db_file, "search", "--product", "apache", "--version", "2.4.49")
    assert result.exit_code == 2
    assert "--vendor" in result.output
    assert "http_server" in result.output
    assert "tomcat" in result.output


def test_search_without_version(db_file: Path) -> None:
    result = _run(db_file, "search", "--product", "httpd")
    assert result.exit_code == 2
    assert "--version" in result.output


def test_search_without_product_or_cpe(db_file: Path) -> None:
    result = _run(db_file, "search", "--version", "1.0")
    assert result.exit_code == 2


# --------------------------------------------------------------------------- #
# resolve / show / stats
# --------------------------------------------------------------------------- #
def test_resolve(db_file: Path) -> None:
    result = _run(db_file, "resolve", "apache http")
    assert result.exit_code == 0
    assert "http_server" in result.stdout


def test_resolve_unknown(db_file: Path) -> None:
    result = _run(db_file, "resolve", "prodotto_inesistente_xyz")
    assert result.exit_code == 2


def test_show(db_file: Path) -> None:
    result = _run(db_file, "show", "CVE-2021-41773")
    assert result.exit_code == 0
    assert "CWE-22" in result.stdout
    assert "sfruttamento noto" in result.stdout  # evidenziazione KEV


def test_show_unknown_cve(db_file: Path) -> None:
    result = _run(db_file, "show", "CVE-1999-0001")
    assert result.exit_code == 2


def test_stats(db_file: Path) -> None:
    result = _run(db_file, "stats")
    assert result.exit_code == 0
    assert "CVE totali" in result.stdout
    assert "apache" in result.stdout


def test_sync_requires_a_mode(db_file: Path) -> None:
    result = _run(db_file, "sync")
    assert result.exit_code == 2
    assert "--full" in result.output


# --------------------------------------------------------------------------- #
# batch
# --------------------------------------------------------------------------- #
def _inventory(path: Path) -> Path:
    path.write_text(
        "host,port,vendor,product,version\n"
        "10.0.0.1,443,apache,http_server,2.4.52\n"
        "10.0.0.2,8080,apache,tomcat,9.0.10\n"
        "10.0.0.3,21,,software_inesistente_xyz,1.0\n",
        encoding="utf-8",
    )
    return path


def test_batch_xlsx(db_file: Path, tmp_path: Path) -> None:
    inventory = _inventory(tmp_path / "inventory.csv")
    output = tmp_path / "risultati.xlsx"
    result = _run(db_file, "batch", "--input", str(inventory), "--output", str(output))

    assert result.exit_code == 1  # trovate CVE >= HIGH
    assert output.exists()

    from openpyxl import load_workbook

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["risultati", "non_risolti"]

    sheet = workbook["risultati"]
    header = [cell.value for cell in sheet[1]]
    assert header[:5] == ["host", "port", "software", "versione", "cve"]
    hosts = {row[0] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert hosts == {"10.0.0.1", "10.0.0.2"}

    unresolved = workbook["non_risolti"]
    rows = list(unresolved.iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "10.0.0.3"
    assert "non presente" in rows[0][5]


def test_batch_csv(db_file: Path, tmp_path: Path) -> None:
    inventory = _inventory(tmp_path / "inventory.csv")
    output = tmp_path / "risultati.csv"
    _run(db_file, "batch", "--input", str(inventory), "--output", str(output))

    assert output.exists()
    assert (tmp_path / "risultati_non_risolti.csv").exists()
    rows = list(csv.DictReader(output.open(encoding="utf-8")))
    assert {row["host"] for row in rows} == {"10.0.0.1", "10.0.0.2"}


def test_batch_missing_input(db_file: Path, tmp_path: Path) -> None:
    result = _run(
        db_file,
        "batch",
        "--input",
        str(tmp_path / "assente.csv"),
        "--output",
        str(tmp_path / "out.xlsx"),
    )
    assert result.exit_code == 2


def test_read_inventory_requires_columns(tmp_path: Path) -> None:
    path = tmp_path / "bad.csv"
    path.write_text("host,port\n1.2.3.4,80\n", encoding="utf-8")
    with pytest.raises(ValueError) as excinfo:
        read_inventory(path)
    assert "product" in str(excinfo.value)


def test_read_inventory_handles_optional_vendor(tmp_path: Path) -> None:
    path = tmp_path / "inv.csv"
    path.write_text("product,version\nhttp_server,2.4.52\n", encoding="utf-8")
    rows = read_inventory(path)
    assert rows[0].vendor is None
    assert rows[0].product == "http_server"


def test_version_flag() -> None:
    result = runner.invoke(app, ["--version"])
    assert result.exit_code == 0
    assert "nvdlocal" in result.stdout


def test_enum_options_are_case_insensitive(db_file: Path) -> None:
    """--min-severity e --format accettano anche minuscolo."""
    result = _run(
        db_file,
        "search",
        "--product",
        "httpd",
        "--version",
        "2.4.49",
        "--min-severity",
        "critical",
        "--format",
        "json",
    )
    payload = json.loads(result.stdout)
    assert {r["cve_id"] for r in payload["results"]} == {"CVE-2022-22720"}


def test_no_subcommand_shows_help(db_file: Path) -> None:
    """Con e senza opzioni globali il comportamento dev'essere lo stesso."""
    bare = runner.invoke(app, [])
    with_option = _run(db_file)
    assert bare.exit_code == with_option.exit_code == 2
    assert "Usage" in bare.output
    assert "Usage" in with_option.output
